import io

import docx
import pymupdf

from app.ingest.chunk import chunk_pages
from app.ingest.extract import extract
from tests.fakes import FakeLLM


def make_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Indexing", fontsize=22)
    page.insert_text((72, 120), "A B-tree index keeps keys sorted for fast lookups.", fontsize=11)
    doc.new_page()  # blank page -> treated as scanned -> OCR
    return doc.tobytes()


def test_pdf_keeps_pages_headings_and_ocrs_blank_pages():
    pages = extract("notes.pdf", make_pdf(), FakeLLM())
    assert [p.page for p in pages] == [1, 2]
    assert pages[0].text.startswith("#") and "B-tree" in pages[0].text
    assert pages[1].text == "OCR TEXT from image"
    chunks = chunk_pages(pages)
    assert chunks[0].heading == "Indexing"


def test_docx_headings_become_markdown():
    document = docx.Document()
    document.add_heading("Transactions", level=1)
    document.add_paragraph("ACID stands for atomicity, consistency, isolation, durability.")
    buffer = io.BytesIO()
    document.save(buffer)
    pages = extract("notes.docx", buffer.getvalue(), None)
    assert "# Transactions" in pages[0].text and "ACID" in pages[0].text


def test_image_uses_ocr():
    assert extract("slide.png", b"\x89PNG", FakeLLM())[0].text == "OCR TEXT from image"
