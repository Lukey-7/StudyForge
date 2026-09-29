"""Step 1 of ingestion: bytes -> list of pages of plain text.

| File type          | Tool                         | Page numbers kept? |
|--------------------|------------------------------|--------------------|
| PDF (digital)      | pymupdf4llm -> markdown      | yes                |
| PDF (scanned page) | render page -> Gemini OCR    | yes                |
| DOCX               | python-docx paragraphs       | no (DOCX has none) |
| TXT / MD           | decode UTF-8                 | no                 |
| Image              | Gemini vision OCR            | page 1             |
| Audio              | Gemini transcription         | no                 |
"""

import io
import logging
from dataclasses import dataclass
from pathlib import PurePath

import docx
import pymupdf
import pymupdf4llm

from app.llm.base import LLM

logger = logging.getLogger(__name__)

OCR_INSTRUCTION = (
    "Transcribe ALL readable text in this image exactly, preserving headings, lists and "
    "reading order. Output plain text only, no commentary. If there is no text, output nothing."
)
AUDIO_INSTRUCTION = (
    "Transcribe this audio recording verbatim in English. Split into paragraphs at topic "
    "changes. Output plain text only, no timestamps or commentary."
)

FILE_TYPES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
    ".mp3": "audio",
    ".wav": "audio",
    ".m4a": "audio",
    ".ogg": "audio",
    ".flac": "audio",
}
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
AUDIO_MIME = {".mp3": "audio/mp3", ".wav": "audio/wav", ".m4a": "audio/mp4", ".ogg": "audio/ogg", ".flac": "audio/flac"}

# A PDF page with fewer characters than this is treated as a scanned image.
SCANNED_PAGE_MIN_CHARS = 25


@dataclass
class ExtractedPage:
    page: int | None  # 1-based page number, None when the format has no pages
    text: str


class UnsupportedFileType(ValueError):
    pass


def detect_file_type(file_name: str) -> str:
    ext = PurePath(file_name).suffix.lower()
    if ext not in FILE_TYPES:
        raise UnsupportedFileType(f"unsupported file type {ext!r}; allowed: {', '.join(sorted(FILE_TYPES))}")
    return FILE_TYPES[ext]


def extract(file_name: str, data: bytes, llm: LLM | None, max_ocr_pages: int = 30) -> list[ExtractedPage]:
    """Dispatch on file type. `llm` is only needed for OCR / audio."""
    file_type = detect_file_type(file_name)
    ext = PurePath(file_name).suffix.lower()
    if file_type == "pdf":
        return extract_pdf(data, llm, max_ocr_pages)
    if file_type == "docx":
        return extract_docx(data)
    if file_type == "text":
        return [ExtractedPage(page=None, text=data.decode("utf-8", errors="replace"))]
    if llm is None:
        raise ValueError(f"{file_type} files need an LLM for OCR/transcription")
    if file_type == "image":
        return [ExtractedPage(page=1, text=llm.read_media(data, IMAGE_MIME[ext], OCR_INSTRUCTION).text)]
    return [ExtractedPage(page=None, text=llm.read_media(data, AUDIO_MIME[ext], AUDIO_INSTRUCTION).text)]


def extract_pdf(data: bytes, llm: LLM | None, max_ocr_pages: int) -> list[ExtractedPage]:
    """pymupdf4llm turns each page into markdown: headings become '#' lines, lists and
    tables are kept, which lets the chunker split by section. Pages with (almost) no
    text are scanned images, so we render them and let Gemini OCR them."""
    pages: list[ExtractedPage] = []
    ocr_used = 0
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        markdown_pages = pymupdf4llm.to_markdown(doc, page_chunks=True)
        for index, page_md in enumerate(markdown_pages):
            text = _clean_markdown(page_md["text"])
            if len(text) < SCANNED_PAGE_MIN_CHARS and llm is not None and ocr_used < max_ocr_pages:
                png = doc[index].get_pixmap(dpi=150).tobytes("png")
                text = llm.read_media(png, "image/png", OCR_INSTRUCTION).text.strip()
                ocr_used += 1
            pages.append(ExtractedPage(page=index + 1, text=text))
    if ocr_used:
        logger.info("PDF OCR used on %d page(s)", ocr_used)
    return pages


def _clean_markdown(text: str) -> str:
    """Drop pymupdf4llm's image placeholders ("==> picture ... omitted <==")."""
    lines = [line for line in text.splitlines() if "intentionally omitted" not in line]
    return "\n".join(lines).strip()


def extract_docx(data: bytes) -> list[ExtractedPage]:
    """Headings become markdown '#' lines so the chunker can detect sections."""
    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading") or style == "title":
            level = style.replace("heading", "").strip()
            hashes = "#" * (int(level) if level.isdigit() else 1)
            lines.append(f"\n{hashes} {text}\n")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [ExtractedPage(page=None, text="\n\n".join(lines))]
